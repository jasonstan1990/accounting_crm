
import os, re, csv, json, sqlite3, unicodedata, os.path as p
from datetime import datetime
from typing import List, Optional, Dict, Any
from zipfile import ZipFile

import fitz  # PyMuPDF
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from openpyxl import Workbook

# ========== Config & Credentials ==========
BASE_DIR = p.dirname(__file__)
VISION_JSON_PATH = p.join(BASE_DIR, "vision.json")
if p.exists(VISION_JSON_PATH):
    os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = VISION_JSON_PATH

DOCAI_CONFIG_PATH = p.join(BASE_DIR, "docai_config.json")
if p.exists(DOCAI_CONFIG_PATH):
    with open(DOCAI_CONFIG_PATH, "r", encoding="utf-8") as f:
        DOCAI_CFG = json.load(f)
else:
    DOCAI_CFG = {"provider":"vision","mode":"auto","project_id":"","location":"eu","processor_id_invoice":"","processor_id_expense":""}

# ========= SQLite =========
DB_PATH = p.join(BASE_DIR, "data.sqlite3")

def db_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def init_db():
    conn = db_conn(); c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS customers(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT,
            afm  TEXT UNIQUE,
            email TEXT,
            phone TEXT,
            notes TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS invoices(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT,
            issuer_afm TEXT,
            customer_afm TEXT,
            issue_date TEXT,
            net_amount REAL,
            vat_amount REAL,
            total_amount REAL,
            vat_rate REAL,
            currency TEXT,
            doc_type TEXT,
            series TEXT,
            number TEXT,
            description TEXT,
            status TEXT,
            raw_text TEXT,
            customer_id INTEGER,
            ocr_confidence REAL,
            field_conf TEXT,
            FOREIGN KEY(customer_id) REFERENCES customers(id) ON DELETE SET NULL
        );
    """)
    # indexes
    c.execute("CREATE INDEX IF NOT EXISTS idx_customers_afm ON customers(afm);")
    c.execute("CREATE INDEX IF NOT EXISTS idx_invoices_cust ON invoices(customer_id);")
    c.execute("CREATE INDEX IF NOT EXISTS idx_invoices_date ON invoices(issue_date);")
    conn.commit(); conn.close()

init_db()

# ========= Optional libs =========
try:
    from google.cloud import vision
    VISION_AVAILABLE = True
except Exception:
    VISION_AVAILABLE = False

try:
    from google.cloud import documentai as docai
    DOCAI_AVAILABLE = True
except Exception:
    DOCAI_AVAILABLE = False

# ========= OCR helpers =========
def run_vision_ocr_with_conf(image_bytes: bytes):
    if not VISION_AVAILABLE:
        raise RuntimeError("google-cloud-vision not installed or missing credentials (vision.json)")
    client = vision.ImageAnnotatorClient()
    resp = client.document_text_detection(image=vision.Image(content=image_bytes))
    if resp.error.message:
        raise RuntimeError(resp.error.message)
    full_text = resp.full_text_annotation.text if resp.full_text_annotation else ""
    confs = []
    if resp.full_text_annotation:
        for page in resp.full_text_annotation.pages:
            for block in page.blocks:
                for para in block.paragraphs:
                    for word in para.words:
                        confs.append(word.confidence or 0.0)
    avg_conf = (sum(confs)/len(confs)) if confs else 0.0
    return full_text, avg_conf

def run_docai_ocr(raw: bytes, mime_type: str = "application/pdf", which: str = "auto"):
    if not DOCAI_AVAILABLE:
        raise RuntimeError("google-cloud-documentai not installed")
    if not DOCAI_CFG.get("project_id") or not DOCAI_CFG.get("location"):
        raise RuntimeError("docai_config.json missing project_id/location")
    location = DOCAI_CFG.get("location")
    project_id = DOCAI_CFG.get("project_id")
    mode = DOCAI_CFG.get("mode") or "auto"
    if which == "auto":
        which = mode

    processor_id = DOCAI_CFG.get("processor_id_invoice") if which == "invoice" else DOCAI_CFG.get("processor_id_expense") or DOCAI_CFG.get("processor_id_invoice")
    if not processor_id:
        raise RuntimeError("docai_config.json missing processor_id")

    # Regional endpoint
    opts = dict(api_endpoint=f"{location}-documentai.googleapis.com")
    client = docai.DocumentProcessorServiceClient(client_options=opts)
    name = client.processor_path(project_id, location, processor_id)

    raw_document = docai.RawDocument(content=raw, mime_type=mime_type)
    request = docai.ProcessRequest(name=name, raw_document=raw_document)
    result = client.process_document(request=request)
    document = result.document

    # Full text
    full_text = document.text or ""
    # Average conf (page level if exists)
    confs = []
    try:
        for pge in document.pages:
            if hasattr(pge, "confidence") and pge.confidence is not None:
                confs.append(float(pge.confidence))
    except Exception:
        pass
    avg_conf = (sum(confs)/len(confs)) if confs else 0.0

    # Map some common fields from entities
    fields = {}
    try:
        for ent in document.entities:
            t = (ent.type_ or "").lower()
            val = ent.mention_text or ent.normalized_value.text if ent.normalized_value else ent.mention_text
            if not val:
                continue
            fields[t] = val
    except Exception:
        pass

    return full_text, avg_conf, fields

def pdf_to_images(pdf_bytes: bytes, dpi: int = 240) -> List[bytes]:
    imgs = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    for pg in doc:
        pix = pg.get_pixmap(dpi=dpi, alpha=False)
        imgs.append(pix.tobytes("png"))
    doc.close()
    return imgs

# ========= Text parsing =========
AFM_TOKEN    = re.compile(r"[AΑ][FΦ][MΜ]", re.IGNORECASE)
AFM_INLINE   = re.compile(r"(\D|^)(\d[\D]*\d[\D]*\d[\D]*\d[\D]*\d[\D]*\d[\D]*\d[\D]*\d[\D]*\d)(\D|$)")
AFM_FALLBACK = re.compile(r"\b(\d[\d\s\.\-]{0,12}\d)\b")
DATE_REGEXES = [re.compile(r"\b(\d{2}[\/\-\.]\d{2}[\/\-\.]\d{4})\b"),
                re.compile(r"\b(\d{4}[\/\-\.]\d{2}[\/\-\.]\d{2})\b")]
MONEY = re.compile(r"(?:€|\bEUR\b)?\s*([0-9]{1,3}(?:[.,\s][0-9]{3})*(?:[.,][0-9]{2})|[0-9]+(?:[.,][0-9]{2}))")

def norm_num(s:str):
    if not s: return None
    s = s.strip()
    if s.count(",")>0 and s.count(".")==0:
        s = s.replace(" ","").replace(".","").replace(",",".")
    else:
        s = s.replace(" ","").replace(",", "")
    try: return float(s)
    except: return None

def is_valid_afm(afm:str)->bool:
    if not afm or not afm.isdigit() or len(afm)!=9: return False
    d=list(map(int,afm)); chk=d[-1]; s=sum(d[i]*(2**(8-i)) for i in range(8))
    return (s%11)%10==chk

def find_afms(text:str):
    out=[]; lines=text.splitlines()
    def _add(c):
        if is_valid_afm(c) and c not in out: out.append(c)
    for line in lines:
        u=unicodedata.normalize("NFKC", line.upper()).replace(" ","")
        if AFM_TOKEN.search(u):
            m=AFM_INLINE.search(line)
            if m:
                digits=re.sub(r"\D","",m.group(0)); m9=re.search(r"\d{9}",digits)
                if m9: _add(m9.group(0))
        if (("ΑΦΜ" in u) or ("AFM" in u)) and ":" in line:
            part=line.split(":",1)[1]; digits=re.sub(r"\D","",part)
            m9=re.search(r"\d{9}",digits); 
            if m9: _add(m9.group(0))
    if not out:
        for m in AFM_FALLBACK.finditer(text):
            digits=re.sub(r"\D","",m.group(1)); m9=re.search(r"\d{9}",digits)
            if m9: _add(m9.group(0))
    return out[:2]

def find_date(text:str, preferred:str=""):
    # try preferred (Document AI field) first
    if preferred:
        for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%Y/%m/%d","%Y.%m.%d","%d.%m.%Y"):
            try: return datetime.strptime(preferred.strip(), fmt).strftime("%Y-%m-%d")
            except: pass
    for rx in DATE_REGEXES:
        m=rx.search(text)
        if m:
            raw=m.group(1)
            for fmt in ("%d/%m/%Y","%d-%m-%Y","%d.%m.%Y","%Y/%m/%d","%Y-%m-%d","%Y.%m.%d"):
                try: return datetime.strptime(raw,fmt).strftime("%Y-%m-%d")
                except: pass
    return ""

def find_amounts(text:str, hints:Dict[str,str]|None=None):
    hints = hints or {}
    vat=None; total=None; net=None; rate=None

    # hints from DocAI
    if hints.get("total_amount"): total = norm_num(hints.get("total_amount"))
    if hints.get("vat_amount"):   vat   = norm_num(hints.get("vat_amount"))
    if hints.get("net_amount"):   net   = norm_num(hints.get("net_amount"))
    if hints.get("vat_rate"):
        try:
            rate = float(re.sub(r"[^0-9\.]", "", hints.get("vat_rate")))
        except: pass

    lines=text.splitlines(); n=len(lines)
    def grab(i):
        for j in range(i,min(i+4,n)):
            for m in MONEY.finditer(lines[j]):
                v=norm_num(m.group(1)); 
                if v is not None: return v
    for i,l in enumerate(lines):
        u=unicodedata.normalize("NFKC", l.upper()).replace(" ","")
        if vat is None and any(k in u for k in ("ΦΠΑ","VAT")): vat=grab(i)
        if total is None and any(k in u for k in ("ΣΥΝΟΛ","TOTAL","ΠΛΗΡΩΤ")): total=grab(i)
        if net is None and any(k in u for k in ("ΚΑΘΑΡ","NET")): net=grab(i)
    if total is None:
        cand=[norm_num(m.group(1)) for m in MONEY.finditer(text)]
        cand=[c for c in cand if c is not None]
        if cand: total=max(cand)
    if net is not None and total is None and vat is not None: total=round(net+vat,2)
    if net is not None and vat is None and total is not None: vat=round(total-net,2)
    if total is not None and vat is not None and net is None: net=round(total-vat,2)
    if rate is None and net and vat and net>0: rate=round(100*vat/net,2)
    return net or 0.0, vat or 0.0, total or 0.0, rate or 0.0

def parse_invoice_text(text:str, filename:str, doc_conf:float, doc_fields:Dict[str,str]|None=None):
    doc_fields = doc_fields or {}
    afms=find_afms(text)
    issuer=doc_fields.get("supplier_tax_id") or (afms[0] if len(afms)>0 else "")
    cust_afm=doc_fields.get("customer_tax_id") or (afms[1] if len(afms)>1 else "")
    issue=find_date(text, preferred=doc_fields.get("invoice_date") or doc_fields.get("date") or "")
    net, vat, total, rate = find_amounts(text, hints=doc_fields)
    currency="EUR" if ("€" in text or "EUR" in text.upper()) else ""
    doc_type="Τιμολόγιο" if ("ΤΙΜΟΛΟΓ" in text.upper() or "INVOICE" in text.upper()) else "Παραστατικό"
    series=doc_fields.get("invoice_series") or ""
    number=doc_fields.get("invoice_number") or doc_fields.get("receipt_number") or ""
    first=text.strip().splitlines()[0] if text.strip().splitlines() else ""
    desc=first[:200]
    field_conf={"issuer_afm":doc_conf, "customer_afm":doc_conf, "net_amount":doc_conf, "vat_amount":doc_conf, "total_amount":doc_conf}
    cust_id=None
    if cust_afm:
        conn=db_conn(); cur=conn.cursor()
        cur.execute("INSERT OR IGNORE INTO customers(name,afm) VALUES(?,?)",(None, cust_afm))
        cur.execute("SELECT id FROM customers WHERE afm=?",(cust_afm,))
        r=cur.fetchone(); conn.commit(); conn.close()
        if r: cust_id=r["id"]
    return dict(filename=filename, issuer_afm=issuer, customer_afm=cust_afm, issue_date=issue,
                net_amount=net, vat_amount=vat, total_amount=total, vat_rate=rate, currency=currency,
                doc_type=doc_type, series=series, number=number, description=desc, status="ok",
                raw_text=text[:6000], customer_id=cust_id, ocr_confidence=doc_conf,
                field_conf=json.dumps(field_conf))

def run_ocr_auto(filename:str, content:bytes):
    name = filename.lower()
    use_docai = (DOCAI_CFG.get("provider") == "docai")
    # prefer DocAI for pdf; for images both work
    if use_docai and DOCAI_AVAILABLE:
        mime = "application/pdf" if name.endswith(".pdf") else ("image/jpeg" if name.endswith((".jpg",".jpeg")) else "image/png")
        try:
            full_text, conf, fields = run_docai_ocr(content, mime_type=mime, which=DOCAI_CFG.get("mode","auto"))
            return full_text, conf, fields
        except Exception as e:
            # fall back to vision if installed
            if VISION_AVAILABLE:
                if name.endswith(".pdf"):
                    pages = pdf_to_images(content)
                    texts, confs = [], []
                    for img in pages:
                        t,c = run_vision_ocr_with_conf(img); texts.append(t); confs.append(c)
                    return "\n\n".join(texts), (sum(confs)/len(confs)) if confs else 0.0, {}
                else:
                    t,c = run_vision_ocr_with_conf(content); return t,c,{}
            raise
    else:
        if not VISION_AVAILABLE:
            raise RuntimeError("Vision OCR not available (install google-cloud-vision and provide vision.json)")
        if name.endswith(".pdf"):
            pages = pdf_to_images(content)
            texts, confs = [], []
            for img in pages:
                t,c = run_vision_ocr_with_conf(img); texts.append(t); confs.append(c)
            return "\n\n".join(texts), (sum(confs)/len(confs)) if confs else 0.0, {}
        else:
            t,c = run_vision_ocr_with_conf(content); return t,c,{}

# ========= Query helpers =========
def rows_to_list(cur): return [dict(r) for r in cur.fetchall()]

def list_invoices(filters:Dict[str,Any], page:int, per_page:int):
    sql="SELECT * FROM invoices WHERE 1=1"; args=[]
    if filters.get("customer_id"): sql+=" AND customer_id=?"; args.append(filters["customer_id"])
    if filters.get("afm"): sql+=" AND (customer_afm=? OR issuer_afm=?)"; args += [filters["afm"], filters["afm"]]
    if filters.get("q"):
        q=f"%{filters['q']}%"; sql+=" AND (filename LIKE ? OR description LIKE ? OR series LIKE ? OR number LIKE ?)"
        args += [q,q,q,q]
    if filters.get("date_from"): sql+=" AND ifnull(issue_date,'') >= ?"; args.append(filters["date_from"])
    if filters.get("date_to"):   sql+=" AND ifnull(issue_date,'') <= ?"; args.append(filters["date_to"])
    sql+=" ORDER BY ifnull(issue_date,'' ) DESC, id DESC LIMIT ? OFFSET ?"
    args += [per_page, (page-1)*per_page]
    conn=db_conn(); cur=conn.cursor(); cur.execute(sql,args); data=rows_to_list(cur)
    cur.execute("SELECT COUNT(*) as c FROM ("+sql.replace(" LIMIT ? OFFSET ?","")+")", args[:-2])
    total=cur.fetchone()["c"]; conn.close()
    return data, total

def get_invoice(iid:int)->Optional[dict]:
    conn=db_conn(); cur=conn.cursor(); cur.execute("SELECT * FROM invoices WHERE id=?", (iid,))
    r=cur.fetchone(); conn.close(); return dict(r) if r else None

def update_invoice(iid:int, upd:dict):
    allowed={'issuer_afm','customer_afm','issue_date','net_amount','vat_amount',
             'total_amount','vat_rate','currency','doc_type','series','number',
             'description','status','customer_id'}
    keys=[k for k in upd.keys() if k in allowed]
    if not keys: return
    sets=", ".join([f"{k}=?" for k in keys]); vals=[upd[k] for k in keys]+[iid]
    conn=db_conn(); cur=conn.cursor(); cur.execute(f"UPDATE invoices SET {sets} WHERE id=?", vals)
    conn.commit(); conn.close()

def create_customer(d:dict)->int:
    conn=db_conn(); cur=conn.cursor()
    cur.execute("INSERT INTO customers(name,afm,email,phone,notes) VALUES(?,?,?,?,?)",
                (d.get("name"), d.get("afm"), d.get("email"), d.get("phone"), d.get("notes")))
    cid=cur.lastrowid; conn.commit(); conn.close(); return cid

def patch_customer(cid:int, d:dict):
    allowed={"name","afm","email","phone","notes"}
    keys=[k for k in d.keys() if k in allowed]
    if not keys: return
    sets=", ".join([f"{k}=?" for k in keys]); vals=[d[k] for k in keys]+[cid]
    conn=db_conn(); cur=conn.cursor(); cur.execute(f"UPDATE customers SET {sets} WHERE id=?", vals)
    conn.commit(); conn.close()

def customer_by_id(cid:int)->Optional[dict]:
    conn=db_conn(); cur=conn.cursor(); cur.execute("SELECT * FROM customers WHERE id=?", (cid,))
    r=cur.fetchone(); conn.close(); return dict(r) if r else None

def search_customers(q: str, page: int, per_page: int):
    conn = db_conn(); cur = conn.cursor()
    sql = """
      SELECT c.*,
             COUNT(i.id)               AS inv_count,
             IFNULL(SUM(i.total_amount),0) AS sum_total
      FROM customers c
      LEFT JOIN invoices i ON i.customer_id = c.id
      WHERE 1=1
    """
    params = []
    if q:
        like = f"%{q}%"
        sql += " AND (IFNULL(c.name,'') LIKE ? OR IFNULL(c.afm,'') LIKE ?"
        params += [like, like]
        norm = re.sub(r"\D", "", q)
        if norm:
            sql += " OR REPLACE(REPLACE(REPLACE(IFNULL(c.afm,''),' ',''),'-',''),'.','') LIKE ?"
            params.append(f"%{norm}%")
            if len(norm) == 9:
                sql += " OR REPLACE(REPLACE(REPLACE(IFNULL(c.afm,''),' ',''),'-',''),'.','') = ?"
                params.append(norm)
        sql += ")"
    sql += " GROUP BY c.id ORDER BY c.created_at DESC LIMIT ? OFFSET ?"
    params += [per_page, (page-1)*per_page]
    cur.execute(sql, params)
    data = [dict(r) for r in cur.fetchall()]

    count_sql = "SELECT COUNT(*) AS c FROM customers WHERE 1=1"
    count_params = []
    if q:
        like = f"%{q}%"
        count_sql += " AND (IFNULL(name,'') LIKE ? OR IFNULL(afm,'') LIKE ?"
        count_params += [like, like]
        norm = re.sub(r"\D", "", q)
        if norm:
            count_sql += " OR REPLACE(REPLACE(REPLACE(IFNULL(afm,''),' ',''),'-',''),'.','') LIKE ?"
            count_params.append(f"%{norm}%")
            if len(norm) == 9:
                count_sql += " OR REPLACE(REPLACE(REPLACE(IFNULL(afm,''),' ',''),'-',''),'.','') = ?"
                count_params.append(norm)
        count_sql += ")"
    cur.execute(count_sql, count_params)
    total = cur.fetchone()["c"]
    conn.close()
    return data, total

def list_customers_with_counts():
    conn=db_conn(); cur=conn.cursor()
    cur.execute("""SELECT c.*,
                   COUNT(i.id) as inv_count,
                   IFNULL(SUM(i.net_amount),0) as sum_net,
                   IFNULL(SUM(i.vat_amount),0) as sum_vat,
                   IFNULL(SUM(i.total_amount),0) as sum_total
                   FROM customers c
                   LEFT JOIN invoices i ON i.customer_id=c.id
                   GROUP BY c.id
                   ORDER BY c.created_at DESC""")
    rows=rows_to_list(cur); conn.close(); return rows

def delete_customer(cid:int, cascade:bool):
    conn=db_conn(); cur=conn.cursor()
    if cascade: cur.execute("DELETE FROM invoices WHERE customer_id=?", (cid,))
    else:       cur.execute("UPDATE invoices SET customer_id=NULL WHERE customer_id=?", (cid,))
    cur.execute("DELETE FROM customers WHERE id=?", (cid,))
    conn.commit(); conn.close()

# ========= FastAPI =========
app = FastAPI(title="AI myDATA Automation — CRM")

app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"]
)

static_path = p.join(BASE_DIR, "static")
app.mount("/static", StaticFiles(directory=static_path), name="static")

@app.get("/", response_class=HTMLResponse)
def home():
    with open(p.join(static_path,"customers.html"), "r", encoding="utf-8") as f:
        return HTMLResponse(f.read())

# ----- Schemas -----
class InvoiceUpdate(BaseModel):
    issuer_afm: Optional[str]=None
    customer_afm: Optional[str]=None
    issue_date: Optional[str]=None
    net_amount: Optional[float]=None
    vat_amount: Optional[float]=None
    total_amount: Optional[float]=None
    vat_rate: Optional[float]=None
    currency: Optional[str]=None
    doc_type: Optional[str]=None
    series: Optional[str]=None
    number: Optional[str]=None
    description: Optional[str]=None
    status: Optional[str]=None
    customer_id: Optional[int]=None

class CustomerCreate(BaseModel):
    name:str
    afm: Optional[str]=None
    email: Optional[str]=None
    phone: Optional[str]=None
    notes: Optional[str]=None

class CustomerPatch(BaseModel):
    name: Optional[str]=None
    afm: Optional[str]=None
    email: Optional[str]=None
    phone: Optional[str]=None
    notes: Optional[str]=None

class AssignPayload(BaseModel):
    ids: List[int]
    customer_id: Optional[int]  # None -> unassign

# ----- Uploads per customer -----
@app.post("/api/customers/{cid}/upload")
async def upload_for_customer(cid: int, files: List[UploadFile] = File(...)):
    cust = customer_by_id(cid)
    if not cust: raise HTTPException(404, "Customer not found")
    results = []
    for f in files:
        content = await f.read()
        try:
            full_text, doc_conf, fields = run_ocr_auto(f.filename, content)
            parsed = parse_invoice_text(full_text, f.filename, doc_conf, fields)
        except Exception as e:
            parsed = dict(filename=f.filename, issuer_afm="", customer_afm="", issue_date="",
                          net_amount=0.0, vat_amount=0.0, total_amount=0.0, vat_rate=0.0,
                          currency="", doc_type="", series="", number="", description="",
                          status=f"error: {e}", raw_text="", customer_id=None,
                          ocr_confidence=0.0, field_conf=json.dumps({}))
        parsed["customer_id"] = cid
        if not parsed.get("customer_afm") and (cust.get("afm") or ""):
            parsed["customer_afm"] = cust["afm"]
        cols = list(parsed.keys()); vals = [parsed[k] for k in cols]
        conn = db_conn(); cur = conn.cursor()
        cur.execute(f"INSERT INTO invoices({','.join(cols)}) VALUES ({','.join(['?']*len(cols))})", vals)
        parsed["id"] = cur.lastrowid
        conn.commit(); conn.close()
        results.append(parsed)
    return {"ok": True, "count": len(results), "invoices": results}

# ----- Invoices API -----
@app.get("/api/invoices")
def api_invoices(q:str="", afm:str="", customer_id:int=0,
                 date_from:str="", date_to:str="", page:int=1, per_page:int=200):
    data,total = list_invoices({"q":q,"afm":afm,"customer_id":customer_id or None,
                                "date_from":date_from or None, "date_to":date_to or None},
                               page=max(1,page), per_page=max(1,min(per_page,500)))
    return {"invoices": data, "total": total, "page": page, "per_page": per_page}

@app.post("/api/invoices/{iid}")
def api_update_invoice(iid:int, payload:InvoiceUpdate):
    current=get_invoice(iid)
    if not current: raise HTTPException(404, "Invoice not found")
    upd={k:v for k,v in payload.dict().items() if v is not None}
    r={**current, **upd}
    net=float(r.get("net_amount") or 0.0)
    vat=float(r.get("vat_amount") or 0.0)
    tot=float(r.get("total_amount") or 0.0)
    rate=float(r.get("vat_rate") or 0.0)
    if "net_amount" in upd and "vat_rate" in upd:
        vat=round(upd["net_amount"]*upd["vat_rate"]/100,2); tot=round(upd["net_amount"]+vat,2)
        upd["vat_amount"]=vat; upd["total_amount"]=tot
    elif "net_amount" in upd and ("vat_amount" in upd or rate):
        if "vat_amount" not in upd: vat=round(upd["net_amount"]*rate/100,2); upd["vat_amount"]=vat
        upd["total_amount"]=round(upd["net_amount"]+upd["vat_amount"],2)
    elif "vat_amount" in upd and ("net_amount" in upd or net):
        base=upd.get("net_amount", net); upd["total_amount"]=round(base+upd["vat_amount"],2)
        upd["vat_rate"]=round(100*upd["vat_amount"]/base,2) if base else 0.0
    elif "total_amount" in upd and ("net_amount" in upd or net):
        base=upd.get("net_amount", net); vat=round(upd["total_amount"]-base,2)
        upd["vat_amount"]=vat; upd["vat_rate"]=round(100*vat/base,2) if base else 0.0
    elif "vat_rate" in upd and ("net_amount" in upd or net):
        base=upd.get("net_amount", net); vat=round(base*upd["vat_rate"]/100,2)
        upd["vat_amount"]=vat; upd["total_amount"]=round(base+vat,2)
    update_invoice(iid, upd)
    return {"ok": True, "invoice": get_invoice(iid)}

@app.post("/api/invoices/assign")
def api_assign(payload: AssignPayload):
    if not payload.ids: raise HTTPException(400, "ids required")
    conn=db_conn(); cur=conn.cursor()
    cur.executemany("UPDATE invoices SET customer_id=? WHERE id=?", [(payload.customer_id, iid) for iid in payload.ids])
    conn.commit(); conn.close()
    return {"ok": True, "updated": len(payload.ids)}

@app.delete("/api/invoices/{iid}")
def api_delete_invoice(iid:int):
    conn=db_conn(); cur=conn.cursor(); cur.execute("DELETE FROM invoices WHERE id=?", (iid,))
    conn.commit(); conn.close(); return {"ok": True, "deleted": iid}

# ----- Customers API -----
@app.get("/api/customers")
def api_customers(q:str="", page:int=1, per_page:int=50):
    data,total = search_customers(q, max(1,page), max(1,min(per_page,200)))
    return {"customers": data, "total": total, "page": page, "per_page": per_page}

@app.get("/api/customers/list")
def api_customers_list():
    return {"customers": list_customers_with_counts()}

@app.post("/api/customers")
def api_create_customer(payload: CustomerCreate):
    cid = create_customer(payload.dict())
    return {"ok": True, "id": cid, "customer": customer_by_id(cid)}

@app.patch("/api/customers/{cid}")
def api_patch_customer(cid:int, payload:CustomerPatch):
    if not customer_by_id(cid): raise HTTPException(404, "Customer not found")
    patch_customer(cid, payload.dict(exclude_none=True)); 
    return {"ok": True, "customer": customer_by_id(cid)}

@app.get("/api/customers/{cid}")
def api_get_customer(cid:int):
    c = customer_by_id(cid)
    if not c: raise HTTPException(404, "Customer not found")
    return {"customer": c}

@app.get("/api/customers/{cid}/overview")
def api_customer_overview(cid:int):
    c = customer_by_id(cid)
    if not c: raise HTTPException(404, "Customer not found")
    invs,_ = list_invoices({"customer_id":cid},1,100000)
    k=dict(net=round(sum((i.get("net_amount") or 0) for i in invs),2),
           vat=round(sum((i.get("vat_amount") or 0) for i in invs),2),
           total=round(sum((i.get("total_amount") or 0) for i in invs),2),
           count=len(invs))
    by={}
    for r in invs:
        m=(r.get("issue_date") or "")[:7] or "Unknown"
        a=by.setdefault(m, {"net":0,"vat":0,"total":0,"count":0})
        a["net"]+=r.get("net_amount") or 0; a["vat"]+=r.get("vat_amount") or 0
        a["total"]+=r.get("total_amount") or 0; a["count"]+=1
    series=[{"month":k, **v} for k,v in sorted(by.items())]
    return {"customer": c, "kpis": k, "series": series, "invoices": invs}

@app.delete("/api/customers/{cid}")
def api_delete_customer(cid:int, cascade:bool=False):
    if not customer_by_id(cid): raise HTTPException(404, "Customer not found")
    delete_customer(cid, cascade); return {"ok": True, "deleted": cid, "cascade": cascade}

# ----- Export -----
@app.get("/api/customers/{cid}/export/xlsx")
def export_xlsx_customer(cid: int):
    c = customer_by_id(cid)
    if not c: raise HTTPException(404, "Customer not found")
    invs, _ = list_invoices({"customer_id": cid}, page=1, per_page=100000)
    if not invs: raise HTTPException(400, "No data")
    wb = Workbook()
    ws = wb.active; ws.title = "Invoices"
    cols = ["id","filename","issue_date","net_amount","vat_amount","total_amount","vat_rate",
            "series","number","description","issuer_afm","customer_afm"]
    ws.append(cols)
    for r in invs: ws.append([r.get(cn, "") for cn in cols])
    ws2 = wb.create_sheet("By Month")
    agg = {}
    for r in invs:
        m = (r.get("issue_date") or "")[:7] or "Unknown"
        a = agg.setdefault(m, {"net":0,"vat":0,"total":0})
        a["net"] += r.get("net_amount") or 0
        a["vat"] += r.get("vat_amount") or 0
        a["total"] += r.get("total_amount") or 0
    ws2.append(["Month","Net","VAT","Total"])
    for k,v in sorted(agg.items()):
        ws2.append([k, round(v["net"],2), round(v["vat"],2), round(v["total"],2)])
    out = os.path.join(BASE_DIR, f"customer_{cid}_invoices.xlsx")
    wb.save(out)
    return FileResponse(out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"customer_{cid}_invoices.xlsx")
